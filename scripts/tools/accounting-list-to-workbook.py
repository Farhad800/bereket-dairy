"""Accountant's list (.xls) → Bereket items workbook (same layout the system downloads), costs and counts empty."""
import sys, re, collections, xlrd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
SRC, FINAL, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
clean = lambda s: re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip()
sh = xlrd.open_workbook(SRC).sheet_by_index(0)
rows = [[clean(sh.cell_value(r, c)) for c in range(4)] for r in range(1, sh.nrows)]
fin = {r[0]: r for r in list(openpyxl.load_workbook(FINAL, read_only=True)["Products"].iter_rows(values_only=True))[1:] if r[0]}
KMAP = {"product": "PRODUCT", "sticker": "STICKER", "material": "MATERIAL", "packaging": "PACKAGING"}
# Codes Excel showed as "2090 4" / "2090!0" are 2090.4 / 2090.10 in the system (same names).
def fix(code):
    m = re.fullmatch(r"(\d+) (\d)", code) or None
    if m: return f"{m[1]}.{m[2]}"
    m = re.fullmatch(r"(\d+)!(\d)", code)
    if m: return f"{m[1]}.1{m[2]}"
    return code
fixed = []
for r in rows:
    c = fix(r[0])
    if c != r[0]:
        assert c in fin and clean(fin[c][1]) == r[1], (r, c)
        fixed.append((r[0], c)); r[0] = c
def kind(r):
    if r[0] in fin: return KMAP[fin[r[0]][3]]
    n = r[1] + " " + r[2]
    if re.search(r"اتيكيت|ETIKET", n, re.I): return "STICKER"
    if re.search(r"نايلون|كارتون|ستريج|شريط|NAYLON|KOLI", n, re.I): return "PACKAGING"
    return "MATERIAL" if r[3] == "KG" else "PRODUCT"
unit = lambda u: "kg" if u.upper() == "KG" else "pcs"

SUP = [("Code", "الرمز", 14, 0), ("Name *", "الاسم", 40, 0), ("Second name", "الاسم الثاني", 34, 0), ("Unit (pcs/kg)", "الوحدة", 12, 0), ("Weight of one bag (kg)", "وزن الكيس الواحد", 16, 0),
       ("Cost price", "سعر التكلفة", 13, 0), ("Currency (IQD/USD/EUR)", "العملة", 14, 0), ("Rate to IQD", "سعر الصرف إلى الدينار", 14, 0), ("Cost in IQD (calculated)", "التكلفة بالدينار (محسوبة)", 16, 1),
       ("In stock now (system)", "الموجود حالياً في النظام", 16, 1), ("Counted stock", "الجرد الفعلي", 14, 0)]
PRO = [("Product code", "رمز المنتج", 16, 0), ("Name *", "الاسم", 40, 0), ("Second name", "الاسم الثاني", 34, 0), ("Mold code", "رمز القالب", 14, 0), ("Label code", "رمز الملصق", 14, 0),
       ("QC required (yes/no)", "فحص الجودة", 12, 0), ("In stock now (system)", "الموجود حالياً في النظام", 16, 1), ("Counted stock", "الجرد الفعلي", 14, 0)]
REC = [("Product code *", "رمز المنتج", 16, 0), ("Product name", "اسم المنتج", 30, 1), ("Item code *", "رمز المادة", 14, 0), ("Item name", "اسم المادة", 30, 1), ("Quantity *", "الكمية", 12, 0), ("Per (piece/pack) *", "لكل (قطعة/عبوة)", 14, 0)]
wb = openpyxl.Workbook(); help_ = wb.active; help_.title = "How to use"; help_.column_dimensions["A"].width = 120
for i, l in enumerate([
    "Bereket Plastic — items workbook (labels, raw materials, packaging, products)", "",
    "1. Every item from the accounting list is here, on its sheet, with its code and names. Fill in the Cost price, Currency, Rate (USD/EUR) and Counted stock.",
    "2. Each row is matched on its Code. Empty cells never delete anything. Grey columns are for information only and are not read back.",
    "3. Cost: price of ONE unit (1 piece or 1 kg) + its currency. For USD or EUR also type today's rate (1 USD = how many IQD).",
    "4. Weight of one bag: only for packaging stocked in kg but used by the bag (nylon).",
    "5. Counted stock: what is really on the shelf now. Empty = the stock stays as it is.",
    "6. Products: Mold code and Label code can be filled in too. Made of & packaging: one row per item a product uses (Per = piece → grams / pieces for ONE piece; Per = pack → how many products go into ONE carton / bag).",
    "7. Upload: Warehouse Stock → Items workbook (Excel) → choose this file → Check file → Save everything.", "",
    "١. كل مواد قائمة المحاسبة موجودة هنا في أوراقها مع الرمز والأسماء. املأ سعر التكلفة والعملة وسعر الصرف (للدولار/اليورو) والجرد الفعلي.",
    "٢. كل صف يُطابَق برمزه. الخلايا الفارغة لا تحذف شيئاً. الأعمدة الرمادية للمعلومات فقط.",
    "٣. التكلفة: سعر الوحدة الواحدة (قطعة أو كيلو) مع العملة. للدولار أو اليورو اكتب سعر الصرف اليوم.",
    "٤. الجرد الفعلي: الكمية الموجودة فعلاً على الرف. فارغ = يبقى المخزون كما هو.",
    "٥. الرفع: مخزون المستودع ← ملف المواد (إكسل) ← اختر الملف ← فحص الملف ← حفظ الكل.",
]):
    c = help_.cell(i + 1, 1, l); c.alignment = Alignment(wrap_text=True)
    if i == 0: c.font = Font(bold=True, size=14)
BLUE, GREY, INFO = PatternFill("solid", fgColor="FF294680"), PatternFill("solid", fgColor="FF6B7280"), PatternFill("solid", fgColor="FFF1F3F6")
def sheet(title, cols, data):
    ws = wb.create_sheet(title); ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 32
    for i, (en, ar, w, info) in enumerate(cols, 1):
        c = ws.cell(1, i, f"{en} · {ar}"); c.font = Font(bold=True, color="FFFFFFFF"); c.fill = GREY if info else BLUE; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(i)].width = w
    for r, vals in enumerate(data, 2):
        for i, v in enumerate(vals, 1):
            if v not in (None, ""): ws.cell(r, i, v).number_format = "@" if i == 1 else "General"
    last = len(data) + 300
    names = [c[0].replace(" *", "") for c in cols]
    for i, (_, _, _, info) in enumerate(cols, 1):
        if info:
            for r in range(2, last + 1): ws.cell(r, i).fill = INFO
    if "Cost in IQD (calculated)" in names:
        p, cu, rt, ci = (L(names.index(n) + 1) for n in ["Cost price", "Currency (IQD/USD/EUR)", "Rate to IQD", "Cost in IQD (calculated)"])
        for r in range(2, last + 1): ws[f"{ci}{r}"] = f'=IF({p}{r}="","",IF(OR({cu}{r}="",{cu}{r}="IQD"),{p}{r},IF({rt}{r}="","rate?",{p}{r}*{rt}{r})))'
    for name, opts in [("Unit (pcs/kg)", "pcs,kg"), ("Currency (IQD/USD/EUR)", "IQD,USD,EUR"), ("QC required (yes/no)", "yes,no"), ("Per (piece/pack)", "piece,pack")]:
        if name in names:
            dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True, showErrorMessage=True, errorTitle="Bereket", error="Choose: " + opts.replace(",", " / "))
            ws.add_data_validation(dv); col = L(names.index(name) + 1); dv.add(f"{col}2:{col}{last}")
    # codes stay text (0007.22, 000001)
    for r in range(2, last + 1): ws.cell(r, 1).number_format = "@"
    return ws
by = collections.defaultdict(list)
for r in rows: by[kind(r)].append(r)
for k, title in [("STICKER", "Labels"), ("MATERIAL", "Raw materials"), ("PACKAGING", "Packaging")]:
    sheet(title, SUP, [[r[0], r[1], r[2] or None, unit(r[3])] for r in by[k]])
sheet("Products", PRO, [[r[0], r[1], r[2] or None] for r in by["PRODUCT"]])
sheet("Made of & packaging", REC, [])
wb.save(OUT)
print({k: len(v) for k, v in by.items()}, "codes fixed:", fixed)
