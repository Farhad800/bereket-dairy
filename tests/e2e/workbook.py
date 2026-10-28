"""The accountant's items workbook: download every item, fill it in with Excel, upload it back (27 Oct 2026).
EMPTY database. Fresh bereket_accept on :3001 (see TESTING.md), then: python3 tests/e2e/workbook.py"""
import io, os, re, subprocess, sys, tempfile
import openpyxl
from playwright.sync_api import sync_playwright
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from stockdoc import stock_doc
B=os.environ.get("BASE","http://localhost:3001"); DB=os.environ.get("PGDB","bereket_accept"); TOKEN=os.environ.get("SETUP_TOKEN","acceptance-setup-token-2026")
PW="OwnerPass2026!"
res=[]
def ck(n,c): res.append(bool(c)); print(("PASS " if c else "FAIL ")+n, flush=True)
def sql(q): return subprocess.run(["su","postgres","-c",f"psql -d {DB} -tAc \"{q}\""],capture_output=True,text=True).stdout.strip()
Q=lambda s: s.replace('"','\\"')
P=Q('"Product"')
def login(pg,u,pw):
    pg.context.clear_cookies(); pg.goto(B+"/login"); pg.fill("#username",u); pg.fill("#password",pw); pg.click("button[type=submit]")
    pg.wait_for_url(re.compile(r".*/(dashboard|change-password)$"), wait_until="commit"); pg.wait_for_load_state("networkidle")
def stock(code): return float(sql(f"SELECT coalesce(sum(qty),0) FROM {Q('\"StockMovement\"')} m JOIN {P} p ON p.id=m.{Q('\"productId\"')} WHERE p.code='{code}'"))
def col(ws, name):
    for c in ws[1]:
        if c.value and str(c.value).split(" ·")[0].replace(" *","").strip()==name: return c.column
    return None
def row_of(ws, code):
    for r in range(2, ws.max_row+1):
        if ws.cell(r,1).value==code: return r
def first_empty(ws):
    r=2
    while any(ws.cell(r,c).value not in (None,"") for c in range(1,4)): r+=1
    return r
TMP=tempfile.mkdtemp()

with sync_playwright() as p:
    br=p.chromium.launch(); errs=[]
    def ctx(**kw):
        pg=br.new_context(**({"viewport":{"width":1366,"height":900},"accept_downloads":True}|kw)).new_page(); pg.on("pageerror", lambda e: errs.append(str(e))); return pg
    A=ctx()
    A.goto(B+"/setup"); A.fill("#token",TOKEN); A.fill("#employeeId","0001"); A.fill("#fullName","Factory Owner"); A.fill("#username","owner")
    A.fill("#password",PW); A.fill("#confirm",PW); A.click("button[type=submit]"); A.wait_for_url("**/dashboard", wait_until="commit")
    pw={}
    for eid,name,user,role,dept in [("1001","Sara Sales","sara","Sales Manager","Sales & Planning"),("1005","Walid Warehouse","walid","Warehouse Team","Warehouse")]:
        A.goto(B+"/admin?tab=employees"); A.click("text=New employee"); A.fill("#employeeId",eid); A.fill("#fullName",name); A.fill("#username",user)
        A.select_option("#roleId", label=role); A.select_option("#departmentId", label=dept); A.click("form button[type=submit]")
        A.wait_for_selector("text=Temporary password"); temp=A.locator("code").nth(1).inner_text()
        pg=ctx(); login(pg,user,temp); pw[user]=f"{user.capitalize()}Factory2026"
        pg.fill("#current",temp); pg.fill("#next",pw[user]); pg.fill("#confirm",pw[user]); pg.click("button[type=submit]"); pg.wait_for_url("**/dashboard", wait_until="commit"); pg.context.close()
    # what is already in the system
    sql(f"INSERT INTO {Q('\"Mold\"')} (id,code,name,{Q('\"updatedAt\"')}) VALUES ('tmold1','MD-BOWL','Bowl mold',now())")
    for i,code,name,kind,unit in [("tlbl1","LBL-0001","Zakho label","STICKER","pcs"),("tpp1","PP-01","PP granules","MATERIAL","kg"),("tcart1","CART-880","Carton 880","PACKAGING","pcs"),("tnyl1","NYL-01","Nylon bag","PACKAGING","kg"),("tprod1","P-100","Bereket bowl","PRODUCT","pcs")]:
        sql(f"INSERT INTO {P} (id,code,name,kind,unit,{Q('\"updatedAt\"')}) VALUES ('{i}','{code}','{name}','{kind}','{unit}',now())")
    sql(f"UPDATE {P} SET {Q('\"labelId\"')}='tlbl1', {Q('\"moldId\"')}='tmold1' WHERE id='tprod1'")
    sql(f"INSERT INTO {Q('\"ProductComponent\"')} (id,{Q('\"productId\"')},{Q('\"componentId\"')},qty,per) VALUES ('rc1','tprod1','tpp1',50,'PIECE'),('rc2','tprod1','tcart1',880,'CARTON')")
    stock_doc(A, B, "INPUT", [("PP-01",100,None),("CART-880",20,None),("P-100",500,None)], reason="COUNT")

    # ── 1. download: every item, with what is known, its stock and cost columns
    A.goto(B+"/inventory"); ck("Warehouse has an Items workbook button", A.get_by_role("link", name="Items workbook (Excel)").count()==1)
    A.get_by_role("link", name="Items workbook (Excel)").click(); A.wait_for_load_state("networkidle")
    with A.expect_download() as d: A.get_by_role("link", name="Download items workbook").click()
    f1=os.path.join(TMP,"items.xlsx"); d.value.save_as(f1)
    wb=openpyxl.load_workbook(f1)
    ck(f"sheets: How to use, Labels, Raw materials, Packaging, Products, Made of & packaging ({wb.sheetnames})", wb.sheetnames==["How to use","Labels","Raw materials","Packaging","Products","Made of & packaging"])
    L,M,K,PR,R=(wb[n] for n in ["Labels","Raw materials","Packaging","Products","Made of & packaging"])
    ck("items already in the system are listed on their sheet", row_of(L,"LBL-0001") and row_of(M,"PP-01") and row_of(K,"CART-880") and row_of(K,"NYL-01") and row_of(PR,"P-100"))
    ck("cost, currency, rate and counted-stock columns are there", all(col(M,c) for c in ["Cost price","Currency (IQD/USD/EUR)","Rate to IQD","Counted stock","In stock now (system)"]))
    ck("stock now is filled in (PP 100 kg, product 500)", M.cell(row_of(M,"PP-01"), col(M,"In stock now (system)")).value==100 and PR.cell(row_of(PR,"P-100"), col(PR,"In stock now (system)")).value==500)
    ck("product's mold and label are filled in", PR.cell(row_of(PR,"P-100"), col(PR,"Mold code")).value=="MD-BOWL" and PR.cell(row_of(PR,"P-100"), col(PR,"Label code")).value=="LBL-0001")
    ck("existing made of & packaging rows are listed", R.max_row>=3 and {(R.cell(r,1).value,R.cell(r,3).value) for r in range(2,R.max_row+1) if R.cell(r,1).value}=={("P-100","PP-01"),("P-100","CART-880")})

    def upload(path, save=False):
        A.goto(B+"/import/items"); A.wait_for_load_state("networkidle")
        A.set_input_files("input[name=file]", path); A.get_by_role("button", name="Check file").click()
        A.wait_for_selector("[data-testid=iw-extra], [data-testid=iw-error]", timeout=60000); A.wait_for_timeout(300)
        if save and A.get_by_role("button", name="Save everything").count():
            A.get_by_role("button", name="Save everything").click(); A.wait_for_selector("[data-testid=iw-done]", timeout=120000)
        return A.inner_text("main")
    out=upload(f1); ck("the file as downloaded changes nothing", "Nothing new or changed" in out or "nothing" in out.lower())

    # ── 2. the accountant fills it in
    r=row_of(M,"PP-01"); M.cell(r,col(M,"Cost price")).value=1.5; M.cell(r,col(M,"Currency (IQD/USD/EUR)")).value="USD"; M.cell(r,col(M,"Rate to IQD")).value=2000; M.cell(r,col(M,"Counted stock")).value=96.5
    r=row_of(K,"CART-880"); K.cell(r,col(K,"Cost price")).value=1500; K.cell(r,col(K,"Counted stock")).value=25
    r=row_of(K,"NYL-01"); K.cell(r,col(K,"Weight of one bag (kg)")).value=0.067; K.cell(r,col(K,"Cost price")).value=0.5; K.cell(r,col(K,"Currency (IQD/USD/EUR)")).value="USD"; K.cell(r,col(K,"Rate to IQD")).value=1500; K.cell(r,col(K,"Counted stock")).value=3.25
    r=row_of(L,"LBL-0001"); L.cell(r,col(L,"Cost price")).value=25
    r=first_empty(L); L.cell(r,col(L,"Name")).value="Duhok label"; L.cell(r,col(L,"Cost price")).value=30; L.cell(r,col(L,"Counted stock")).value=1000
    r=first_empty(M); M.cell(r,col(M,"Code")).value="dye-blk"; M.cell(r,col(M,"Name")).value="Black dye"; M.cell(r,col(M,"Unit (pcs/kg)")).value="kg"; M.cell(r,col(M,"Cost price")).value=12000
    r=first_empty(PR); PR.cell(r,col(PR,"Product code")).value="P-200"; PR.cell(r,col(PR,"Name")).value="Duhok bowl"; PR.cell(r,col(PR,"Mold code")).value="MD-BOWL"; PR.cell(r,col(PR,"Label code")).value="LBL-0001"; PR.cell(r,col(PR,"QC required (yes/no)")).value="yes"
    r=row_of(PR,"P-100"); PR.cell(r,col(PR,"Counted stock")).value=480
    nr=first_empty(R)
    for i,(prod,item,qty,per) in enumerate([("P-200","PP-01",48,"piece"),("P-200","DYE-BLK",0.5,"piece"),("P-200","CART-880",880,"pack"),("P-200","NYL-01",880,"pack")]):
        for c,v in zip([1,3,5,6],[prod,item,qty,per]): R.cell(nr+i,c).value=v
    # a label row in a product's block = that product's label; an empty row of a block is skipped
    r=first_empty(L); L.cell(r,col(L,"Code")).value="LBL-X1"; L.cell(r,col(L,"Name")).value="Erbil label"
    R.cell(nr+4,1).value="P-100"; R.cell(nr+4,3).value="LBL-X1"; R.cell(nr+4,6).value="piece"
    R.cell(nr+5,1).value="P-100"; R.cell(nr+5,6).value="pack"
    f2=os.path.join(TMP,"filled.xlsx"); wb.save(f2)
    # a copy with mistakes
    bad=openpyxl.load_workbook(f2); bM=bad["Raw materials"]; bR=bad["Made of & packaging"]; bK=bad["Packaging"]; bP=bad["Products"]
    rr=row_of(bM,"DYE-BLK") or first_empty(bM); bM.cell(row_of(bM,"PP-01"),col(bM,"Rate to IQD")).value=None
    bK.cell(row_of(bK,"CART-880"),col(bK,"Counted stock")).value=2.5
    r=first_empty(bR); bR.cell(r,1).value="P-200"; bR.cell(r,3).value="PP-01"; bR.cell(r,5).value=10; bR.cell(r,6).value="pack"
    r=first_empty(bP); bP.cell(r,col(bP,"Name")).value="No code no mold"
    f3=os.path.join(TMP,"bad.xlsx"); bad.save(f3)
    before=sql(f"SELECT count(*) FROM {P}")
    out=upload(f3); errs_txt=A.get_by_test_id("iw-errors").inner_text() if A.get_by_test_id("iw-errors").count() else ""
    ck("mistakes are listed by sheet and row, nothing saved", all(x in errs_txt for x in ["Type the rate: 1 USD","whole number","isn't packaging","needs a Product code or a Mold code"]) and A.get_by_role("button", name="Save everything").count()==0 and sql(f"SELECT count(*) FROM {P}")==before)

    out=upload(f2)
    plan=A.get_by_test_id("iw-plan").inner_text().split("\n")
    ck(f"check shows what will change ({' | '.join(plan[1:])})", "Costs set or changed: 6" in out and "Stock count: 3 items go up, 2 go down" in out)
    A.get_by_role("button", name="Save everything").click(); A.wait_for_selector("[data-testid=iw-done]", timeout=120000); out=A.inner_text("main")
    yr=sql("SELECT extract(year from now())::int")
    ck("saved, with one Input and one Output count document", f"IN-{yr}-0002" in out and f"OUT-{yr}-0001" in out)

    # ── 3. what is in the system now
    cost=lambda c: sql(f"SELECT {Q('\"costPrice\"')}::float||' '||{Q('\"costCurrency\"')}||' x'||{Q('\"costRate\"')}::float||'='||{Q('\"costIqd\"')}::float FROM {P} WHERE code='{c}'")
    ck(f"PP cost 1.5 USD × 2,000 = 3,000 IQD ({cost('PP-01')})", cost("PP-01")=="1.5 USD x2000=3000")
    ck(f"nylon 0.5 USD × 1,500 = 750 IQD, bag 0.067 kg ({cost('NYL-01')})", cost("NYL-01")=="0.5 USD x1500=750" and sql(f"SELECT {Q('\"unitWeightKg\"')}::float FROM {P} WHERE code='NYL-01'")=="0.067")
    ck("carton 1,500 IQD, label 25 IQD", cost("CART-880")=="1500 IQD x1=1500" and cost("LBL-0001")=="25 IQD x1=25")
    ck("new label got an automatic code LBL-0002 with its cost", sql(f"SELECT code||' '||{Q('\"costIqd\"')}::float FROM {P} WHERE name='Duhok label'")=="LBL-0002 30")
    ck("new raw material DYE-BLK (code in capitals, kg, 12,000 IQD)", sql(f"SELECT kind||' '||unit||' '||{Q('\"costIqd\"')}::float FROM {P} WHERE code='DYE-BLK'")=="MATERIAL kg 12000")
    ck("new product P-200 on the mold with label and QC", sql(f"SELECT p.kind||' '||m.code||' '||l.code||' '||p.{Q('\"qcRequired\"')} FROM {P} p JOIN {Q('\"Mold\"')} m ON m.id=p.{Q('\"moldId\"')} JOIN {P} l ON l.id=p.{Q('\"labelId\"')} WHERE p.code='P-200'")=="PRODUCT MD-BOWL LBL-0001 true")
    ck("P-200 made of & packaging: 4 lines", sql(f"SELECT string_agg(c.code||':'||pc.qty::float||':'||pc.per,',' ORDER BY c.code) FROM {Q('\"ProductComponent\"')} pc JOIN {P} c ON c.id=pc.{Q('\"componentId\"')} JOIN {P} p ON p.id=pc.{Q('\"productId\"')} WHERE p.code='P-200'")=="CART-880:880:CARTON,DYE-BLK:0.5:PIECE,NYL-01:880:CARTON,PP-01:48:PIECE")
    ck("P-100's made of & packaging kept as it was", sql(f"SELECT count(*) FROM {Q('\"ProductComponent\"')} WHERE {Q('\"productId\"')}='tprod1'")=="2")
    ck("a label listed under P-100 became its label; the empty row was skipped", sql(f"SELECT l.code FROM {P} p JOIN {P} l ON l.id=p.{Q('\"labelId\"')} WHERE p.code='P-100'")=="LBL-X1")
    s={c: stock(c) for c in ["PP-01","CART-880","NYL-01","LBL-0002","P-100"]}
    ck(f"stock set to the counted amounts: PP 96.5, cartons 25, nylon 3.25, Duhok label 1,000, product 480 ({s})", s=={"PP-01":96.5,"CART-880":25,"NYL-01":3.25,"LBL-0002":1000,"P-100":480})
    ck("count documents have reason COUNT and the file name", sql(f"SELECT string_agg(type||':'||reason,',' ORDER BY type) FROM {Q('\"StockDoc\"')} WHERE note LIKE 'Stock count from Excel (filled.xlsx)'")=="INPUT:COUNT,OUTPUT:COUNT")
    A.goto(B+"/inventory?tab=packaging"); ck("warehouse shows the new values (cartons 25 × 1,500 = 37,500)", A.get_by_test_id("value-CART-880").inner_text()=="37,500")
    out=upload(f2); ck("uploading the same file again changes nothing (new rows found by name, counts already match)", "Stock count: 0 items go up, 0 go down" in out and "Costs set or changed: 0" in out and A.get_by_role("button", name="Save everything").count()==0 and sql(f"SELECT count(*) FROM {P} WHERE name='Duhok label'")=="1")
    ck("audit log has the download and the upload", {"items.workbookDownload","items.import"}<=set(sql(f"SELECT string_agg(DISTINCT action, ',') FROM {Q('\"AuditLog\"')} WHERE action LIKE 'items.%'").split(",")))

    # ── 4. who may use it
    W=ctx(); login(W,"walid",pw["walid"])
    ck("warehouse can't download it", W.request.get(B+"/api/import/items").status==403)
    W.goto(B+"/import/items"); ck("…nor open the page", "/forbidden" in W.url)
    S=ctx(); login(S,"sara",pw["sara"])
    r=S.request.get(B+"/api/import/items"); ws=openpyxl.load_workbook(io.BytesIO(r.body()))["Raw materials"]
    ck("Sales (no money) gets the workbook without cost columns", r.status==200 and col(ws,"Cost price") is None and col(ws,"Name"))
    S.goto(B+"/import/items"); S.set_input_files("input[name=file]", f2); S.get_by_role("button", name="Check file").click(); S.wait_for_selector("[data-testid=iw-extra], [data-testid=iw-error]", timeout=60000)
    ck("…and can't upload costs", "permission to set costs" in S.inner_text("main"))

    # ── 5. phone + Arabic
    Mo=ctx(**p.devices["iPhone 13"]); login(Mo,"owner",PW); Mo.goto(B+"/import/items"); Mo.wait_for_load_state("networkidle")
    ck("phone layout fits", Mo.evaluate("document.documentElement.scrollWidth - window.innerWidth")<=1)
    A.context.add_cookies([{"name":"bfms_locale","value":"ar","url":B}]); A.goto(B+"/import/items"); ck("Arabic page", "ملف المواد" in A.inner_text("main"))
    ck("no browser errors", not errs); print("   errors:", errs[:3])
    br.close()
print(f"\n{sum(res)}/{len(res)} passed"); sys.exit(0 if all(res) else 1)
